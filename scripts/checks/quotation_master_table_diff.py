#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""報價單彙整總表 vs 資料庫（weekly 120，2026-09-07）。

owner：「是否部分案件無匯入？例如『苗栗大山南台鐵平交道潛鑽工程地形測量』系統查詢不到。」

實查：那一案在總表的「系統報價單」第 81 列，**已成立、28,000 元**，
而它**沒有報價單編號** —— 而匯入是以報價單編號為鍵 ⇒ 這種列必然被略過，
**而且不會有任何錯誤訊息**（略過與匯入成功在畫面上長得一樣）。

同型共 4 筆（含 2,186,100 元那筆），另有 6 筆有編號卻不在資料庫裡
（首版說 8 筆，其中 2 筆實際是**同一案的不同版次**——`B114-C031-0` 對上資料庫的
`-1`——逐字比對把版次差異報成缺件。判準已分開，那 2 筆列為 YELLOW）。

## 為什麼做成檢核而不是一次補完

補完是一次性的；下次總表再多幾列沒有編號，同樣會安靜地漏掉。
這支回答的是「**現在總表與系統差在哪裡**」，隨時可跑：

| 判定 | 意思 |
|---|---|
| **RED** | 總表有編號、資料庫沒有 ⇒ 漏匯入 |
| **RED** | 總表列沒有報價單編號 ⇒ 匯入的鍵不存在，必然漏掉（且不會報錯） |
| **RED** | 總表未標成立、系統卻已成案 ⇒ 兩邊不一致（owner 09-07：「誤植已成案刪除，
  系統對應機制請檢視與調整」）。**匯入是單向的**：一列標成立會長出 PM 案 → 承攬案 →
  自動應收，而總表改回未成立或刪掉，系統這邊不會回退、也沒有任何訊號 |
| YELLOW | 資料庫有、總表沒有 ⇒ 附下游影響（成案／請款／發票／金額），
  用來分辨「總表只是沒收錄」與「誤植成案且已產生金流」 |

⚠️ 只比對「系統報價單」這一張工作表：它是總表自己宣告的彙整表
（其餘是各承辦的作業表與備份，內容會重複）。

用法：`python scripts/checks/quotation_master_table_diff.py [xlsx 路徑]`
預設路徑取 `QUOTATION_MASTER_XLSX` 環境變數，再退回 `D:/報價單/115報價單彙整總表.xlsx`。
檔案不存在時回 YELLOW（**不是 GREEN** —— 「沒看到」不等於「沒問題」）。
"""
from __future__ import annotations

import json
import os
from decimal import Decimal
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib.paths import repo_root  # noqa: E402
from lib.docker_exec import python_in  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = repo_root()
DEFAULT_XLSX = os.environ.get("QUOTATION_MASTER_XLSX", r"D:/報價單/115報價單彙整總表.xlsx")
SHEET = "系統報價單"


def read_master(path: Path) -> tuple[dict[str, dict], list[dict]]:
    """回 ({報價單編號: 該列}, [沒有編號的列])。"""
    import openpyxl  # 延後匯入：沒有 openpyxl 時只影響這一支

    ws = openpyxl.load_workbook(path, data_only=True)[SHEET]

    # ⚠️ 2026-09-08：表頭**不在第 1 列**。「系統報價單」第 1 列是給人看的欄位註記
    # （`total_price` / `tax_amount` / `grand_total`），真正的表頭在第 2 列。
    # 首版讀 `ws[1]` ⇒ `col()` 一律找不到名字、全部退回位置 fallback ——
    # 目前碰巧對，但**欄位一移動就會安靜地讀錯欄**（而它不會報錯）。
    # ⇒ 改成找「哪一列有『報價單編號』」，不假設列號。
    hdr_row = 1
    for rn in range(1, 6):
        vals = [str(c.value).strip() if c.value else "" for c in ws[rn]]
        if "報價單編號" in vals:
            hdr_row = rn
            break
    hdr = [str(c.value).strip() if c.value else "" for c in ws[hdr_row]]

    def col(name: str, fallback: int) -> int:
        return hdr.index(name) if name in hdr else fallback

    ci = {k: col(k, i) for i, k in enumerate(
        ("序號", "年度", "承辦同仁", "報價單編號", "是否成立", "報價日期", "客戶名稱", "案名"))}
    ci["報價金額"] = col("報價金額", 9)
    # ⚠️ 同一個表頭在「系統報價單」出現兩次（右側第 38–43 欄是各承辦年度小計），
    #    `col()` 用 `.index()` 取的是**最左邊**那一個 —— 與匯入器 2026-09-07 的修法同一個判準。
    #
    # ⭐ 2026-09-08 owner「經費仍對應不一致」＋「拿總表原檔做一次全欄對帳」：
    # 此前這一支只比對**存在性**（哪些列漏匯入），不比對金額 ——
    # 而當天的每一個問題都在金額欄：總價欄存的是未稅（12 筆）、稅額未填（11 筆）、
    # 佔位發票號未換成真號（32 張）、系統補建的發票金額用了請款額（14 張）。
    # 「有匯入」與「匯對了」是兩個問題，而只驗前者會給出一片綠。
    for k, fb in (("稅內含", 10), ("稅額", 11), ("總價", 12),
                  ("發票號碼", 30), ("銷售額", 31), ("稅額(發票)", 32), ("發票金額", 33)):
        ci[k] = col(k, fb)

    keyed: dict[str, dict] = {}
    unkeyed: list[dict] = []
    for rn, r in enumerate(ws.iter_rows(min_row=hdr_row + 1, values_only=True), start=hdr_row + 1):
        if not any(c for c in r[:10]):
            continue
        row = {k: (str(r[i]).strip() if i < len(r) and r[i] is not None else "")
               for k, i in ci.items()}
        row["_row"] = rn
        if row["報價單編號"]:
            keyed[row["報價單編號"]] = row
        else:
            unkeyed.append(row)
    return keyed, unkeyed


def db_records():
    """{舊編號: {案號, 成案, 請款數, 發票數, 總價}}。

    帶下游影響是為了回答 owner 2026-09-07 的「系統對應機制請檢視與調整」：
    總表一列誤植成「已成立」，匯入會一路長出 PM 案 → 承攬案 →（09-03 起）自動應收；
    而**把總表那一列改掉或刪掉，系統這邊不會回退，也沒有任何訊號**。
    ⇒ 差異要看得到「這一筆已經長出多少東西」，才判得出清理成本。
    """
    sql = (
        "SELECT q.legacy_quotation_no, q.case_code, q.total_price, "
        "EXISTS(SELECT 1 FROM contract_projects c WHERE c.case_code = q.case_code), "
        "(SELECT count(*) FROM erp_billings b WHERE b.erp_quotation_id = q.id), "
        "(SELECT count(*) FROM erp_invoices i WHERE i.erp_quotation_id = q.id), "
        "q.tax_amount, "
        "(SELECT string_agg(i2.invoice_number, ',') FROM erp_invoices i2 WHERE i2.erp_quotation_id = q.id), "
        "(SELECT COALESCE(sum(i3.amount),0) FROM erp_invoices i3 WHERE i3.erp_quotation_id = q.id), "
        "(SELECT COALESCE(sum(i4.tax_amount),0) FROM erp_invoices i4 WHERE i4.erp_quotation_id = q.id) "
        "FROM erp_quotations q "
        "WHERE q.legacy_quotation_no IS NOT NULL AND q.deleted_at IS NULL"
    )
    out = python_in(
        "import asyncio, json\n"
        "from sqlalchemy import text\n"
        "from app.db.database import AsyncSessionLocal\n"
        "SQL = " + repr(sql) + "\n"
        "async def m():\n"
        "    async with AsyncSessionLocal() as db:\n"
        "        rows = (await db.execute(text(SQL))).all()\n"
        "    print('@@' + json.dumps([[r[0], r[1], str(r[2] or ''), bool(r[3]), int(r[4]), int(r[5]), str(r[6] or ''), r[7] or '', str(r[8] or ''), str(r[9] or '')] for r in rows]))\n"
        "asyncio.run(m())\n"
    )
    line = [l for l in (out or "").splitlines() if l.startswith("@@")]
    if not line:
        return None
    return {str(r[0]).strip(): {"case_code": r[1], "total": r[2], "promoted": r[3],
                                "bills": r[4], "invoices": r[5], "tax": r[6],
                                "inv_nos": r[7], "inv_amt": r[8], "inv_tax": r[9]}
            for r in json.loads(line[-1][2:])}


def main() -> int:
    print("=== 報價單彙整總表 vs 資料庫（weekly 120）===")
    path = Path(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX)
    if not path.exists():
        print(f"[YELLOW] 找不到總表 {path} —— 未驗（「沒看到」不等於「沒問題」）")
        print("        可用 QUOTATION_MASTER_XLSX 指定路徑，或以參數傳入。")
        return 1
    try:
        keyed, unkeyed = read_master(path)
    except Exception as exc:  # noqa: BLE001
        print(f"[YELLOW] 讀不了總表：{str(exc)[:120]}")
        return 1

    if not keyed and not unkeyed:
        # 「系統報價單」是**動態陣列公式**（B2 是 ArrayFormula，整張表由 VSTACK 溢出），
        # 而 openpyxl 讀的是 Excel 存檔時的快取值。若有人用不會計算公式的工具存過這個檔，
        # 快取會全空 ⇒ 這裡讀到 0 列。那是「讀不到」，不是「總表是空的」。
        print(f"[YELLOW] 「{SHEET}」讀不到任何資料列 —— 它是動態陣列公式表，"
              f"快取值可能未更新（用 Excel 開啟並存檔即可重算）。未驗。")
        return 1

    recs = db_records()
    if recs is None:
        print("[YELLOW] 連不到資料庫，未驗")
        return 1
    db = set(recs)

    missing_raw = sorted(set(keyed) - db)
    extra = sorted(db - set(keyed))

    # ⚠️ 逐字比對會把「同一案的不同版次」報成漏匯入：實測 `B114-C031-0`（總表）
    # 對上資料庫的 `B114-C031-1`、`B115-022a-1` 對上 `B115-022a-0`。
    # 那是版次差異，不是缺件 —— 直接當成 RED 交出去就是把誤報說成事實。
    # ⇒ 去掉尾碼版次後再比一次，只有**連詞幹都對不到**的才算漏匯入。
    def stem(no: str) -> str:
        base = no.rsplit("-", 1)[0] if "-" in no else no
        return base.lower()

    db_stems = {stem(n) for n in db}
    missing = [n for n in missing_raw if stem(n) not in db_stems]
    version_only = [n for n in missing_raw if stem(n) in db_stems]
    print(f"  總表「{SHEET}」有編號 {len(keyed)} 筆／無編號 {len(unkeyed)} 筆；"
          f"資料庫有舊編號 {len(db)} 筆")

    rc = 0
    if unkeyed:
        print(f"[RED] {len(unkeyed)} 筆**沒有報價單編號** —— 匯入以編號為鍵，"
              f"這些列必然被略過，而且不會有任何錯誤訊息：")
        for r in unkeyed:
            print(f"    第{r['_row']}列 | {r['年度']} | {r['承辦同仁']} | "
                  f"{r['案名'][:34]} | {r['報價金額']}")
        print("      修法：在總表補上報價單編號後重跑匯入（編號要人給，系統不代為編）。")
        rc = 2
    if missing:
        print(f"[RED] {len(missing)} 筆總表有編號、資料庫沒有 ⇒ 漏匯入：")
        for n in missing[:20]:
            r = keyed[n]
            print(f"    {n} | {r['年度']} | {r['承辦同仁']} | {r['案名'][:30]}")
        rc = 2
    if version_only:
        print(f"[YELLOW] {len(version_only)} 筆是**同一案的不同版次**（詞幹在資料庫裡對得到）"
              f"，不算漏匯入，但兩邊的版次不一致：")
        for n in version_only:
            same = sorted(x for x in db if stem(x) == stem(n))
            print(f"    總表 {n} ↔ 資料庫 {'、'.join(same)}")
        rc = max(rc, 1)
    # ③ 成立狀態不一致：總表已改成「未成立」，而系統已經成案。
    #    owner 2026-09-07：「誤植已成案刪除，系統對應機制請檢視與調整」——
    #    匯入是**單向**的：一列誤植成「已成立」會長出 PM 案 → 承攬案 →（09-03 起）自動應收；
    #    把總表那一列改回未成立或刪掉，系統這邊**不會回退，也沒有任何訊號**。
    #    這一條就是那個訊號。判紅是因為它代表帳上多了一筆不該存在的應收。
    established_mismatch = []
    for no, row in keyed.items():
        r = recs.get(no)
        if not r or not r["promoted"]:
            continue
        if str(row.get("是否成立", "")).strip().lower() not in ("v", "y", "yes", "是", "✓", "1", "true"):
            established_mismatch.append((no, row, r))
    if established_mismatch:
        print(f"[RED] {len(established_mismatch)} 筆**總表未標成立、系統卻已成案** —— 兩邊不一致，要人判：")
        print("      · 若總表漏填 → 補上「v」（已開發票的那幾筆多半是這一種）")
        print("      · 若真的誤植成案 → 系統要撤：匯入**不會**自動回退，"
              "成案會一路長出承攬案與自動應收，帳上就多一筆")
        for no, row, r in established_mismatch[:15]:
            print(f"    {no} | {r['case_code']} | 總價 {r['total'] or '-'} | "
                  f"請款 {r['bills']} 張／發票 {r['invoices']} 張 | {row['案名'][:24]}")
        rc = 2

    # ⭐ 2026-09-08：金額欄逐欄對帳（owner「經費仍對應不一致」→「拿總表原檔做一次全欄對帳」）。
    #
    # 此前只比對存在性 ——「有匯入」與「匯對了」是兩個問題，而只驗前者會給出一片綠。
    # 當天在金額欄找到四類，每一類都不會報錯：
    #   · 總價欄存的是未稅（匯入器把總表的「報價金額」＝未稅寫進含稅欄位）12 筆
    #   · 稅額未填（總表有、DB 是 0）11 筆
    #   · 佔位發票號未換成真號 32 張
    #   · 系統補建的發票金額用了**請款額**而不是真實發票金額 14 張
    #
    # ⚠️ 判準要問到每一個欄位：我第一版抽總表時漏了「銷售額」與「稅額(發票)」，
    #    於是對帳報「發票稅額不符 0 筆」——**那不是相符，是根本沒有比對**。
    def _num(v):
        try:
            return Decimal(str(v).replace(",", "").strip())
        except Exception:
            return None

    # 一票多案：同一個發票號出現在總表多列時，總表每列記的是**整張發票**的金額，
    # 而 DB 記的是**該案分攤到的金額** ⇒ 逐列比對必然不符，而兩邊都是對的。
    # （實測 ZX19612010 173,250 分給兩案：137,403 ＋ 35,847 ＝ 173,250）
    # ⇒ 發票欄改成「同號的 DB 金額合計 vs 總表金額」。
    from collections import defaultdict
    inv_rows = defaultdict(list)
    for _no, _row in keyed.items():
        _n = str(_row.get("發票號碼", "")).strip().upper()
        if _n:
            inv_rows[_n].append(_no)

    amt_diff = []
    master_todo = []
    for no, row in keyed.items():
        r = recs.get(no)
        if not r:
            continue
        for label, mkey, dkey in (("總價", "總價", "total"), ("稅額", "稅額", "tax"),
                                  ("發票金額", "發票金額", "inv_amt"), ("發票稅額", "稅額(發票)", "inv_tax")):
            mv, dv = _num(row.get(mkey)), _num(r.get(dkey))
            if mv is None or dv is None:
                continue
            if label.startswith("發票"):
                _n = str(row.get("發票號碼", "")).strip().upper()
                if not _n:
                    continue      # 總表沒開票就不比對發票欄
                peers = inv_rows.get(_n, [no])
                if len(peers) > 1:
                    # 一票多案 ⇒ 合計比對，且只在第一列報一次
                    if no != peers[0]:
                        continue
                    tot = sum((_num(recs[p].get(dkey)) or Decimal(0)) for p in peers if recs.get(p))
                    if abs(mv - tot) > 1:
                        amt_diff.append((no, f"{r['case_code']}（一票多案 {len(peers)} 案）",
                                         label, mv, tot))
                    continue
            if abs(mv - dv) > 1:
                # ⚠️ 方向要分得開：總表那格是空的／0、而 DB 有值且自洽
                # （稅額 ≈ 含稅/21）⇒ 那是**總表待補**，不是系統錯。
                # 判成 RED 會讓人去把對的資料改成錯的。
                _tp = _num(r.get("total"))
                if mv == 0 and dv > 0 and label in ("稅額", "發票稅額") and _tp and _tp > 0                         and abs(dv - (_tp / 21).quantize(Decimal("1"))) <= 2:
                    master_todo.append((no, r["case_code"], label, dv))
                    continue
                amt_diff.append((no, r["case_code"], label, mv, dv))
        mno = str(row.get("發票號碼", "")).strip().upper()
        # 一票多案的 DB 端用 `XLS-<案號>` 佔位＋分攤列表達，不逐案存同一個真號 ⇒ 不判它號碼不符
        if mno and len(inv_rows.get(mno, [])) > 1:
            mno = ""
        if mno and mno not in (r.get("inv_nos") or "").upper():
            amt_diff.append((no, r["case_code"], "發票號碼", mno, (r.get("inv_nos") or "(無)")))
    if amt_diff:
        print(f"[RED] {len(amt_diff)} 處金額／發票欄與總表不符 —— 總表是原件，系統該跟它：")
        for no, cc, label, mv, dv in amt_diff[:20]:
            print(f"    {no} | {cc} | {label}：總表={mv} DB={dv}")
        if len(amt_diff) > 20:
            print(f"    …另 {len(amt_diff) - 20} 處")
        rc = 2
    if master_todo:
        print(f"[YELLOW] {len(master_todo)} 處**總表那一格是空的、而系統有自洽的值** —— 總表待補，"
              f"系統不用改：")
        for no, cc, label, dv in master_todo[:10]:
            print(f"    {no} | {cc} | {label} 總表未填，系統值 {dv}（與總價自洽）")
        rc = max(rc, 1)

    if extra:
        print(f"[YELLOW] {len(extra)} 筆資料庫有、總表沒有 —— 附下游影響，"
              f"用來分辨「總表只是沒收錄」與「誤植成案且已產生金流」：")
        for no in extra[:20]:
            r = recs[no]
            flag = "成案" if r["promoted"] else "未成案"
            print(f"    {no} | {r['case_code']} | {flag} | 總價 {r['total'] or '-'} | "
                  f"請款 {r['bills']}／發票 {r['invoices']}")
        if len(extra) > 20:
            print(f"    …另 {len(extra) - 20} 筆")
        rc = max(rc, 1)

    if rc == 0:
        print("[GREEN] 總表與資料庫一致")
    return rc


if __name__ == "__main__":
    sys.exit(main())
